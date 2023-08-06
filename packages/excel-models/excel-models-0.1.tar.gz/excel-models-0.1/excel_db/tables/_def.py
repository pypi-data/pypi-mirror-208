import typing

from openpyxl.worksheet.worksheet import Worksheet

from ..db import ExcelDB
from ..models import TModel
from ..utils.descriptors import BasePropertyDescriptor


class ExcelTableDefinition(
    BasePropertyDescriptor[ExcelDB],
    typing.Generic[TModel],
):
    title_row: int = 1

    def __init__(
            self,
            model: typing.Type[TModel],
            *,
            table_class: typing.Type['TTable'] = None,
            **kwargs,
    ):
        super().__init__(**kwargs)
        self.model = model
        if table_class is None:
            from ._inst import ExcelTable
            table_class = ExcelTable
        self.table_class: typing.Type['TTable'] = table_class

    def _add_to_class(self):
        self.obj_type.tables.append(self)

    _f_initialize = None

    def _initialize_title_row(self, db: ExcelDB, ws: Worksheet):
        for i, column in enumerate(self.model.columns):
            ws.cell(self.title_row, i + 1, column.name)

    def _initialize_default(self, db: ExcelDB, ws: Worksheet):
        pass

    @property
    def _initialize_method(self):
        if self._f_initialize is None:
            return self._initialize_default
        else:
            return self._f_initialize

    def initialize(self, db: ExcelDB, ws: Worksheet):
        self._initialize_title_row(db, ws)
        self._initialize_method(db, ws)  # noqa: pycharm

    def initializer(self, f_initialize):
        self._f_initialize = f_initialize
        return self

    def _get_default(self, db: ExcelDB) -> Worksheet:
        if self.name in db.wb:
            return db.wb[self.name]
        else:
            ws = db.wb.create_sheet(self.name)
            self.initialize(db, ws)
            return ws

    def _get(self, db: ExcelDB) -> 'TTable':
        if self.attr not in db.ws_cache:
            ws = self._get_method(db)
            db.ws_cache[self.attr] = ws
        ws = db.ws_cache[self.attr]
        return self.table_class(db, self, ws)

    def _set_default(self, db: ExcelDB, ws: Worksheet) -> Worksheet:
        if self.name in db.wb:
            del db.wb[self.name]
        copy: Worksheet = db.wb.copy_worksheet(ws)
        copy.title = self.name
        return copy

    def _set(self, db: ExcelDB, ws: typing.Union[Worksheet, 'TTable']):
        from ._inst import ExcelTable
        if isinstance(ws, ExcelTable):
            ws = ws.ws
        copy = self._set_method(db, ws)
        db.ws_cache[self.attr] = copy

    def _delete_default(self, db: ExcelDB):
        del db.wb[self.name]

    def _delete(self, db: ExcelDB):
        if self.attr in db.__dict__:
            del db.ws_cache[self.attr]
        self._delete_method(db)


TTableDef = typing.TypeVar('TTableDef', bound=ExcelTableDefinition)

if typing.TYPE_CHECKING:
    from ._inst import TTable
